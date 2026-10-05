import io
import re
import zipfile


def extract_text(data: bytes, filename: str) -> str:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "pdf":
        import fitz

        with fitz.open(stream=data, filetype="pdf") as pdf:
            text = "\n\n".join(page.get_text("text") for page in pdf)
    elif suffix == "docx":
        from docx import Document

        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError("El archivo DOCX no es válido")
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    elif suffix == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        text = "\n".join(" | ".join(str(value) for value in row if value is not None) for sheet in workbook for row in sheet.iter_rows(values_only=True))
    elif suffix == "xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data)
        text = "\n".join(" | ".join(str(value) for value in sheet.row_values(row)) for sheet in workbook.sheets() for row in range(sheet.nrows))
    elif suffix == "csv":
        text = data.decode("utf-8-sig", errors="replace")
    else:
        raise ValueError(f"Formato documental no admitido: .{suffix}")
    normalized = re.sub(r"[ \t]+", " ", text.replace("\x00", " "))
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    if not normalized:
        raise ValueError("No se encontró texto legible en el documento")
    return normalized
